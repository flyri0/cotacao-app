import base64
import io
import re
import sqlite3
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Estilos padrão com identidade visual profissional e consistente
HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")  # Azul institucional elegante
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

TITLE_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="0F172A")

SUBTITLE_FONT = Font(name="Calibri", size=9.5, italic=True, color="475569")

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

ROW_EVEN_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
ROW_ODD_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")


def _sanitize_filename(text: str) -> str:
    """Higieniza string para uso seguro como nome de arquivo no Windows/Linux."""
    safe = re.sub(r'[\\/*?:"<>|]', "", text).strip()
    safe = re.sub(r"\s+", "_", safe)
    return safe or "planilha"


def generate_quote_template_excel(
    conn: sqlite3.Connection, id_rodada: int, id_fornecedor: Optional[int] = None
) -> Dict[str, Any]:
    """
    Gera um arquivo Excel (.xlsx) contendo as cotações reais da rodada ativa,
    com 100% de fidelidade aos dados exibidos na tela de Cotações do programa:
    Produto, Categoria, Fornecedor, Marca, Embalagem, Qtd / Emb, Unidade, Preço da Embalagem e Preço Unitário.
    """
    cursor = conn.cursor()

    # Obtém dados da rodada
    cursor.execute("SELECT id, descricao, status, data_criacao FROM rodadas WHERE id = ?", (id_rodada,))
    rodada_row = cursor.fetchone()
    if not rodada_row:
        raise ValueError(f"Rodada #{id_rodada} não encontrada.")
    rodada = dict(rodada_row)

    # Dados do fornecedor se fornecido
    fornecedor_nome = ""
    if id_fornecedor:
        cursor.execute("SELECT nome FROM fornecedores WHERE id = ?", (id_fornecedor,))
        forn_row = cursor.fetchone()
        if forn_row:
            fornecedor_nome = forn_row["nome"]

    # Busca as cotações reais cadastradas na rodada
    query = """
        SELECT 
            c.id,
            c.id_rodada,
            c.id_fornecedor,
            f.nome AS fornecedor_nome,
            c.id_produto,
            p.nome AS produto_nome,
            p.categoria AS produto_categoria,
            c.marca,
            c.embalagem,
            c.qtd_por_embalagem,
            c.unidade,
            c.preco_embalagem,
            (c.preco_embalagem / c.qtd_por_embalagem) AS preco_unitario
        FROM cotacoes c
        JOIN fornecedores f ON f.id = c.id_fornecedor
        JOIN produtos p ON p.id = c.id_produto
        WHERE c.id_rodada = ?
    """
    params: List[Any] = [id_rodada]
    if id_fornecedor:
        query += " AND c.id_fornecedor = ?"
        params.append(id_fornecedor)

    query += " ORDER BY p.categoria ASC, p.nome ASC, f.nome ASC"
    cursor.execute(query, tuple(params))
    cotacoes_existentes = [dict(r) for r in cursor.fetchall()]

    itens_exportacao: List[Dict[str, Any]] = []
    if cotacoes_existentes:
        itens_exportacao = cotacoes_existentes
    else:
        # Fallback: se a rodada ainda não possuir cotações lançadas, lista as necessidades
        cursor.execute(
            """
            SELECT 
                n.id_produto,
                p.nome AS produto_nome,
                p.categoria AS produto_categoria
            FROM necessidades n
            JOIN produtos p ON p.id = n.id_produto
            WHERE n.id_rodada = ?
            ORDER BY p.categoria ASC, p.nome ASC
            """,
            (id_rodada,),
        )
        for r in cursor.fetchall():
            itens_exportacao.append(
                {
                    "id": "-",
                    "produto_nome": r["produto_nome"],
                    "produto_categoria": r["produto_categoria"] or "-",
                    "fornecedor_nome": fornecedor_nome or "-",
                    "marca": "-",
                    "embalagem": "-",
                    "qtd_por_embalagem": 1.0,
                    "unidade": "UN",
                    "preco_embalagem": None,
                    "preco_unitario": None,
                }
            )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotações"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título e Cabeçalho Institucional
    ws.merge_cells("A1:J1")
    ws["A1"] = f"PLANILHA DE COTAÇÕES — {rodada['descricao'].upper()}"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = TITLE_FILL
    ws.row_dimensions[1].height = 32

    subtitulo = f"Rodada #{rodada['id']} | Data de Abertura: {rodada['data_criacao']}"
    if fornecedor_nome:
        subtitulo += f" | Fornecedor: {fornecedor_nome}"
    subtitulo += f" | Total de Itens: {len(itens_exportacao)} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    ws.merge_cells("A2:J2")
    ws["A2"] = subtitulo
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10  # Espaçador sutil

    # 2. Cabeçalho das Colunas (Bate 100% com a visualização do programa, sem Qtd Necessária)
    headers = [
        ("ID", 8),
        ("Produto / Item", 36),
        ("Categoria", 20),
        ("Fornecedor", 28),
        ("Marca", 20),
        ("Descrição da Embalagem", 26),
        ("Qtd na Embalagem", 18),
        ("Unidade Medida", 16),
        ("Preço da Embalagem (R$)", 24),
        ("Preço Unitário (R$)", 22),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_idx, (header_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 3. Linhas de Dados com Cotações Reais
    current_row = 5
    for idx, item in enumerate(itens_exportacao):
        fill = ROW_EVEN_FILL if idx % 2 == 0 else ROW_ODD_FILL

        # Coluna 1: ID
        c_id = ws.cell(row=current_row, column=1, value=item.get("id"))
        c_id.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna 2: Produto
        c_nome = ws.cell(row=current_row, column=2, value=item.get("produto_nome"))
        c_nome.alignment = Alignment(horizontal="left", vertical="center")
        c_nome.font = Font(name="Calibri", size=10, bold=True)

        # Coluna 3: Categoria
        c_cat = ws.cell(row=current_row, column=3, value=item.get("produto_categoria") or "-")
        c_cat.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna 4: Fornecedor
        c_forn = ws.cell(row=current_row, column=4, value=item.get("fornecedor_nome") or "-")
        c_forn.alignment = Alignment(horizontal="left", vertical="center")

        # Coluna 5: Marca Ofertada
        c_marca = ws.cell(row=current_row, column=5, value=item.get("marca") or "-")
        c_marca.alignment = Alignment(horizontal="left", vertical="center")

        # Coluna 6: Embalagem
        c_emb = ws.cell(row=current_row, column=6, value=item.get("embalagem") or "-")
        c_emb.alignment = Alignment(horizontal="left", vertical="center")

        # Coluna 7: Qtd na Embalagem
        qtd_val = item.get("qtd_por_embalagem")
        c_qtd = ws.cell(row=current_row, column=7, value=qtd_val if qtd_val is not None else 1.0)
        c_qtd.alignment = Alignment(horizontal="right", vertical="center")
        c_qtd.number_format = "#,##0.00"

        # Coluna 8: Unidade de Medida
        c_un = ws.cell(row=current_row, column=8, value=item.get("unidade") or "UN")
        c_un.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna 9: Preço da Embalagem
        preco_emb_val = item.get("preco_embalagem")
        c_preco_emb = ws.cell(row=current_row, column=9, value=preco_emb_val if preco_emb_val is not None else "")
        c_preco_emb.alignment = Alignment(horizontal="right", vertical="center")
        c_preco_emb.number_format = "R$ #,##0.00"

        # Coluna 10: Preço Unitário Calculado
        preco_un_val = item.get("preco_unitario")
        c_preco_un = ws.cell(row=current_row, column=10, value=preco_un_val if preco_un_val is not None else "")
        c_preco_un.alignment = Alignment(horizontal="right", vertical="center")
        c_preco_un.font = Font(name="Calibri", size=10, bold=True, color="0F766E")
        c_preco_un.number_format = "R$ #,##0.00"

        # Aplica bordas e preenchimento zebra
        for col_i in range(1, 11):
            cell = ws.cell(row=current_row, column=col_i)
            cell.border = BORDER_THIN
            cell.fill = fill

        ws.row_dimensions[current_row].height = 22
        current_row += 1

    # Congela painéis nos cabeçalhos
    ws.freeze_panes = "A5"

    # Salva o arquivo em memória
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    b64_content = base64.b64encode(buffer.read()).decode("utf-8")

    # Nome do arquivo baseado na descrição real da rodada (sanitizado)
    safe_rodada = _sanitize_filename(rodada["descricao"])
    if fornecedor_nome:
        safe_forn = _sanitize_filename(fornecedor_nome)
        nome_arquivo = f"{safe_rodada}_{safe_forn}.xlsx"
    else:
        nome_arquivo = f"{safe_rodada}.xlsx"

    return {
        "sucesso": True,
        "nome_arquivo": nome_arquivo,
        "conteudo_base64": b64_content,
        "total_itens": len(itens_exportacao),
    }


def process_quote_excel(
    conn: sqlite3.Connection,
    id_rodada: int,
    id_fornecedor: int,
    conteudo_base64: str,
) -> Dict[str, Any]:
    """
    Processa um arquivo Excel (.xlsx) enviado via base64, mapeando cabeçalhos por nome
    (compatível com o modelo novo de 10 colunas e modelos legado de 8/9 colunas).
    """
    raw_data = base64.b64decode(conteudo_base64)
    buffer = io.BytesIO(raw_data)
    wb = openpyxl.load_workbook(buffer, data_only=True)
    ws = wb.active

    cursor = conn.cursor()

    # Valida fornecedor e rodada
    cursor.execute("SELECT id, nome FROM fornecedores WHERE id = ?", (id_fornecedor,))
    forn_row = cursor.fetchone()
    if not forn_row:
        raise ValueError(f"Fornecedor #{id_fornecedor} não encontrado.")

    cursor.execute("SELECT id, status FROM rodadas WHERE id = ?", (id_rodada,))
    rod_row = cursor.fetchone()
    if not rod_row:
        raise ValueError(f"Rodada #{id_rodada} não encontrada.")

    if rod_row["status"] != "aberta":
        raise ValueError(f"Não é permitido importar cotações em rodada {rod_row['status']}.")

    # Carrega mapa de produtos existentes
    cursor.execute("SELECT id, nome FROM produtos")
    produtos_rows = cursor.fetchall()
    produtos_por_id = {r["id"]: r["nome"] for r in produtos_rows}
    produtos_por_nome = {r["nome"].strip().lower(): r["id"] for r in produtos_rows}

    importados = 0
    ignorados = 0
    produtos_criados = 0
    erros: List[str] = []

    # Localiza dinamicamente a linha de cabeçalho e as colunas correspondentes
    header_row_idx = 4
    col_id = 1
    col_produto = 2
    col_categoria = 3
    col_marca = 5
    col_embalagem = 6
    col_qtd = 7
    col_unidade = 8
    col_preco = 9

    for r in range(1, 10):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, 15)]
        has_id = any("id" in v for v in row_vals[:3])
        has_prod = any("produto" in v or "item" in v for v in row_vals[:4])
        if has_id or has_prod:
            header_row_idx = r
            for c in range(1, 15):
                header_str = str(ws.cell(row=r, column=c).value or "").strip().lower()
                if "id" in header_str and not col_id:
                    col_id = c
                elif ("produto" in header_str or "item" in header_str) and not col_produto:
                    col_produto = c
                elif "categoria" in header_str:
                    col_categoria = c
                elif "marca" in header_str:
                    col_marca = c
                elif "embalagem" in header_str and "preço" not in header_str and "preco" not in header_str:
                    if "qtd" in header_str or "quantidade" in header_str:
                        col_qtd = c
                    else:
                        col_embalagem = c
                elif "unidade" in header_str:
                    col_unidade = c
                elif "preço" in header_str or "preco" in header_str:
                    if "unit" not in header_str:
                        col_preco = c
            break

    # Itera pelas linhas de dados
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        id_prod_val = ws.cell(row=row_idx, column=col_id).value if col_id else None
        nome_prod_val = ws.cell(row=row_idx, column=col_produto).value if col_produto else None
        marca_val = ws.cell(row=row_idx, column=col_marca).value if col_marca else None
        embalagem_val = ws.cell(row=row_idx, column=col_embalagem).value if col_embalagem else "Unidade"
        qtd_emb_val = ws.cell(row=row_idx, column=col_qtd).value if col_qtd else 1.0
        unidade_val = ws.cell(row=row_idx, column=col_unidade).value if col_unidade else "UN"
        preco_emb_val = ws.cell(row=row_idx, column=col_preco).value if col_preco else None

        if not id_prod_val and not nome_prod_val:
            continue

        # Valida se há preço preenchido antes de qualquer ação
        if preco_emb_val is None or str(preco_emb_val).strip() == "":
            ignorados += 1
            continue

        try:
            preco_emb = float(preco_emb_val)
        except (ValueError, TypeError):
            erros.append(f"Linha {row_idx}: Preço inválido '{preco_emb_val}'.")
            ignorados += 1
            continue

        if preco_emb <= 0:
            ignorados += 1
            continue

        # Identifica produto existente por ID ou nome
        id_produto = None
        if id_prod_val:
            try:
                id_tentativa = int(id_prod_val)
                if id_tentativa in produtos_por_id:
                    id_produto = id_tentativa
            except (ValueError, TypeError):
                pass

        if not id_produto and nome_prod_val:
            nome_norm = str(nome_prod_val).strip().lower()
            id_produto = produtos_por_nome.get(nome_norm)

        # Auto-cadastro de produto caso não exista e haja preço cotado
        if not id_produto and nome_prod_val and str(nome_prod_val).strip():
            nome_novo = str(nome_prod_val).strip()
            cat_val = ws.cell(row=row_idx, column=col_categoria).value if col_categoria else None
            cat_novo = str(cat_val).strip() if cat_val and str(cat_val).strip() != "-" else "Geral"

            cursor.execute(
                "INSERT INTO produtos (nome, categoria, ativo) VALUES (?, ?, 1)",
                (nome_novo, cat_novo),
            )
            id_produto = cursor.lastrowid
            produtos_por_id[id_produto] = nome_novo
            produtos_por_nome[nome_novo.lower()] = id_produto
            produtos_criados += 1

            # Auto-adiciona na necessidade da rodada
            cursor.execute(
                "INSERT OR IGNORE INTO necessidades (id_rodada, id_produto) VALUES (?, ?)",
                (id_rodada, id_produto),
            )

        if not id_produto:
            erros.append(f"Linha {row_idx}: Produto não identificado e sem nome para criação.")
            ignorados += 1
            continue

        # Sanitiza quantidade e unidade
        try:
            qtd_emb = float(qtd_emb_val) if qtd_emb_val is not None else 1.0
            if qtd_emb <= 0:
                qtd_emb = 1.0
        except (ValueError, TypeError):
            qtd_emb = 1.0

        marca_desc = (
            str(marca_val).strip() if marca_val and str(marca_val).strip() and str(marca_val).strip() != "-" else None
        )
        embalagem_desc = (
            str(embalagem_val).strip() if embalagem_val and str(embalagem_val).strip() != "-" else "Unidade"
        )
        unidade_desc = str(unidade_val).strip().upper() if unidade_val and str(unidade_val).strip() != "-" else "UN"

        # Insere ou atualiza na tabela de cotações
        cursor.execute(
            """
            INSERT INTO cotacoes (
                id_rodada, id_fornecedor, id_produto,
                marca, embalagem, qtd_por_embalagem, unidade, preco_embalagem
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(id_rodada, id_fornecedor, id_produto) DO UPDATE SET
                marca = excluded.marca,
                embalagem = excluded.embalagem,
                qtd_por_embalagem = excluded.qtd_por_embalagem,
                unidade = excluded.unidade,
                preco_embalagem = excluded.preco_embalagem
            """,
            (id_rodada, id_fornecedor, id_produto, marca_desc, embalagem_desc, qtd_emb, unidade_desc, preco_emb),
        )
        importados += 1

    conn.commit()

    return {
        "sucesso": True,
        "importados": importados,
        "produtos_criados": produtos_criados,
        "ignorados": ignorados,
        "erros": erros,
        "fornecedor_nome": forn_row["nome"],
    }


def export_products_excel_db(conn: sqlite3.Connection) -> Dict[str, Any]:
    """Exporta o catálogo completo de produtos com a mesma identidade visual institucional."""
    cursor = conn.cursor()
    cursor.execute("SELECT id, nome, categoria FROM produtos ORDER BY categoria ASC, nome ASC")
    produtos = [dict(r) for r in cursor.fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título e Cabeçalho Institucional
    ws.merge_cells("A1:C1")
    ws["A1"] = "CATÁLOGO GERAL DE PRODUTOS"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = TITLE_FILL
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:C2")
    ws["A2"] = f"Total de Produtos: {len(produtos)} | Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 10  # Espaçador

    # 2. Cabeçalhos de Colunas
    headers = [("ID", 8), ("Nome do Produto", 42), ("Categoria", 26)]
    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_i, (h_text, width) in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_i, value=h_text)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(col_i)].width = width

    # 3. Linhas de Dados com Zebra Striping
    current_row = 5
    for idx, p in enumerate(produtos):
        fill = ROW_EVEN_FILL if idx % 2 == 0 else ROW_ODD_FILL

        c_id = ws.cell(row=current_row, column=1, value=p["id"])
        c_id.alignment = Alignment(horizontal="center", vertical="center")

        c_nome = ws.cell(row=current_row, column=2, value=p["nome"])
        c_nome.alignment = Alignment(horizontal="left", vertical="center")
        c_nome.font = Font(name="Calibri", size=10, bold=True)

        c_cat = ws.cell(row=current_row, column=3, value=p["categoria"] or "-")
        c_cat.alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 4):
            cell = ws.cell(row=current_row, column=c)
            cell.border = BORDER_THIN
            cell.fill = fill

        ws.row_dimensions[current_row].height = 22
        current_row += 1

    ws.freeze_panes = "A5"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    b64 = base64.b64encode(buffer.read()).decode("utf-8")

    return {
        "sucesso": True,
        "nome_arquivo": "catalogo_produtos.xlsx",
        "conteudo_base64": b64,
        "total": len(produtos),
    }


def import_products_excel_db(conn: sqlite3.Connection, conteudo_base64: str) -> Dict[str, Any]:
    """Importa produtos em lote a partir de uma planilha Excel."""
    raw = base64.b64decode(conteudo_base64)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active

    cursor = conn.cursor()
    importados = 0
    ignorados = 0
    erros: List[str] = []

    # Procura linha onde começam os dados
    start_row = 2
    for r in range(1, 6):
        val1 = str(ws.cell(row=r, column=1).value or "").strip().lower()
        val2 = str(ws.cell(row=r, column=2).value or "").strip().lower()
        if "id" in val1 or "produto" in val2 or "nome" in val2:
            start_row = r + 1
            break

    for row_idx in range(start_row, ws.max_row + 1):
        nome_val = ws.cell(row=row_idx, column=2).value
        # Caso a planilha não tenha coluna ID e comece direto no nome
        if not nome_val:
            nome_val = ws.cell(row=row_idx, column=1).value
            cat_val = ws.cell(row=row_idx, column=2).value
        else:
            cat_val = ws.cell(row=row_idx, column=3).value

        if not nome_val or str(nome_val).strip() == "":
            continue

        nome = str(nome_val).strip()
        cat = str(cat_val).strip() if cat_val and str(cat_val).strip() != "-" else None

        try:
            cursor.execute(
                """
                INSERT INTO produtos (nome, categoria, ativo)
                VALUES (?, ?, 1)
                ON CONFLICT(nome) DO UPDATE SET
                    categoria = excluded.categoria,
                    ativo = 1
                """,
                (nome, cat),
            )
            importados += 1
        except Exception as e:
            erros.append(f"Linha {row_idx}: {e}")
            ignorados += 1

    conn.commit()
    return {"sucesso": True, "importados": importados, "ignorados": ignorados, "erros": erros}


def import_suppliers_excel_db(conn: sqlite3.Connection, conteudo_base64: str) -> Dict[str, Any]:
    """Importa fornecedores em lote a partir de uma planilha Excel."""
    raw = base64.b64decode(conteudo_base64)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active

    cursor = conn.cursor()
    importados = 0
    ignorados = 0
    erros: List[str] = []

    start_row = 2
    for r in range(1, 6):
        val1 = str(ws.cell(row=r, column=1).value or "").strip().lower()
        val2 = str(ws.cell(row=r, column=2).value or "").strip().lower()
        if "id" in val1 or "fornecedor" in val2 or "razão" in val2 or "razao" in val2:
            start_row = r + 1
            break

    for row_idx in range(start_row, ws.max_row + 1):
        nome_val = ws.cell(row=row_idx, column=2).value
        if not nome_val:
            nome_val = ws.cell(row=row_idx, column=1).value
            contato_val = ws.cell(row=row_idx, column=2).value
            tel_val = ws.cell(row=row_idx, column=3).value
            email_val = ws.cell(row=row_idx, column=4).value
            min_val = ws.cell(row=row_idx, column=5).value
        else:
            contato_val = ws.cell(row=row_idx, column=3).value
            tel_val = ws.cell(row=row_idx, column=4).value
            email_val = ws.cell(row=row_idx, column=5).value
            min_val = ws.cell(row=row_idx, column=6).value

        if not nome_val or str(nome_val).strip() == "":
            continue

        nome = str(nome_val).strip()
        contato = str(contato_val).strip() if contato_val and str(contato_val).strip() != "-" else None
        tel = str(tel_val).strip() if tel_val and str(tel_val).strip() != "-" else None
        email = str(email_val).strip() if email_val and str(email_val).strip() != "-" else None

        try:
            ped_min = float(min_val) if min_val else 0.0
        except (ValueError, TypeError):
            ped_min = 0.0

        try:
            cursor.execute(
                """
                INSERT INTO fornecedores (nome, contato, telefone, email, pedido_minimo, ativo)
                VALUES (?, ?, ?, ?, ?, 1)
                ON CONFLICT(nome) DO UPDATE SET
                    contato = excluded.contato,
                    telefone = excluded.telefone,
                    email = excluded.email,
                    pedido_minimo = excluded.pedido_minimo,
                    ativo = 1
                """,
                (nome, contato, tel, email, ped_min),
            )
            importados += 1
        except Exception as e:
            erros.append(f"Linha {row_idx}: {e}")
            ignorados += 1

    conn.commit()
    return {"sucesso": True, "importados": importados, "ignorados": ignorados, "erros": erros}
