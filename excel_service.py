import base64
import io
import sqlite3
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Estilos padrão para planilhas profissionais
HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")  # Azul profissional
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

TITLE_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="0F172A")

SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="475569")

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

ROW_EVEN_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
ROW_ODD_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")


def generate_quote_template_excel(
    conn: sqlite3.Connection, id_rodada: int, id_fornecedor: Optional[int] = None
) -> Dict[str, Any]:
    """
    Gera um arquivo Excel (.xlsx) estilizado contendo todos os produtos em necessidade
    da rodada para o fornecedor (ou comprador) preencher os preços e embalagens.
    """
    cursor = conn.cursor()

    # Obtém dados da rodada
    cursor.execute("SELECT id, descricao, status, data_criacao FROM rodadas WHERE id = ?", (id_rodada,))
    rodada_row = cursor.fetchone()
    if not rodada_row:
        raise ValueError(f"Rodada #{id_rodada} não encontrada.")
    rodada = dict(rodada_row)

    # Dados do fornecedor se fornecido
    fornecedor_nome = ""
    if id_fornecedor:
        cursor.execute("SELECT nome FROM fornecedores WHERE id = ?", (id_fornecedor,))
        forn_row = cursor.fetchone()
        if forn_row:
            fornecedor_nome = forn_row["nome"]

    # Busca necessidades da rodada com dados do produto
    cursor.execute(
        """
        SELECT 
            n.id_produto,
            p.nome AS produto_nome,
            p.categoria AS produto_categoria
        FROM necessidades n
        JOIN produtos p ON p.id = n.id_produto
        WHERE n.id_rodada = ?
        ORDER BY p.categoria ASC, p.nome ASC
        """,
        (id_rodada,),
    )
    necessidades = [dict(r) for r in cursor.fetchall()]

    # Busca cotações prévias já lançadas para esse fornecedor nessa rodada se houver
    cotacoes_existentes = {}
    if id_fornecedor:
        cursor.execute(
            """
            SELECT id_produto, marca, embalagem, qtd_por_embalagem, unidade, preco_embalagem
            FROM cotacoes
            WHERE id_rodada = ? AND id_fornecedor = ?
            """,
            (id_rodada, id_fornecedor),
        )
        for row in cursor.fetchall():
            cotacoes_existentes[row["id_produto"]] = dict(row)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotações"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título e Cabeçalho do Documento
    ws.merge_cells("A1:I1")
    ws["A1"] = f"PLANILHA DE COTAÇÃO DE PREÇOS — {rodada['descricao'].upper()}"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = TITLE_FILL
    ws.row_dimensions[1].height = 30

    subtitulo = f"Rodada #{rodada['id']} | Data: {rodada['data_criacao']}"
    if fornecedor_nome:
        subtitulo += f" | Fornecedor: {fornecedor_nome}"
    subtitulo += " | Preencha as colunas destacadas (Marca, Embalagem, Qtd, Unidade e Preço)"

    ws.merge_cells("A2:I2")
    ws["A2"] = subtitulo
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 2. Cabeçalho das Colunas
    headers = [
        ("ID", 8),
        ("Produto / Item", 35),
        ("Categoria", 18),
        ("Qtd Necessária", 15),
        ("Marca Ofertada", 22),
        ("Descrição da Embalagem", 28),
        ("Qtd na Embalagem", 18),
        ("Unidade Medida", 16),
        ("Preço da Embalagem (R$)", 22),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_idx, (header_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 3. Linhas de Dados
    current_row = 5
    for idx, nec in enumerate(necessidades):
        id_prod = nec["id_produto"]
        cot = cotacoes_existentes.get(id_prod, {})

        marca_val = cot.get("marca", "")
        embalagem_val = cot.get("embalagem", "Unidade")
        qtd_emb_val = cot.get("qtd_por_embalagem", 1.0)
        unidade_val = cot.get("unidade", "UN")
        preco_emb_val = cot.get("preco_embalagem", None)

        fill = ROW_EVEN_FILL if idx % 2 == 0 else ROW_ODD_FILL

        # Coluna A: ID Produto
        c_id = ws.cell(row=current_row, column=1, value=id_prod)
        c_id.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna B: Nome do Produto
        c_nome = ws.cell(row=current_row, column=2, value=nec["produto_nome"])
        c_nome.alignment = Alignment(horizontal="left", vertical="center")
        c_nome.font = Font(name="Calibri", size=10, bold=True)

        # Coluna C: Categoria
        c_cat = ws.cell(row=current_row, column=3, value=nec["produto_categoria"] or "-")
        c_cat.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna D: Quantidade Solicitada (opcional/não definida)
        c_qtd = ws.cell(row=current_row, column=4, value=nec.get("quantidade_solicitada") or "-")
        c_qtd.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna E: Marca Ofertada (Preenchimento)
        c_marca = ws.cell(row=current_row, column=5, value=marca_val)
        c_marca.alignment = Alignment(horizontal="left", vertical="center")

        # Coluna F: Descrição da Embalagem (Preenchimento)
        c_emb = ws.cell(row=current_row, column=6, value=embalagem_val)
        c_emb.alignment = Alignment(horizontal="left", vertical="center")

        # Coluna G: Qtd na Embalagem (Preenchimento)
        c_qtd_emb = ws.cell(row=current_row, column=7, value=qtd_emb_val)
        c_qtd_emb.alignment = Alignment(horizontal="right", vertical="center")
        c_qtd_emb.number_format = "#,##0.00"

        # Coluna H: Unidade de Medida (Preenchimento)
        c_un = ws.cell(row=current_row, column=8, value=unidade_val)
        c_un.alignment = Alignment(horizontal="center", vertical="center")

        # Coluna I: Preço da Embalagem (Preenchimento)
        c_preco = ws.cell(row=current_row, column=9, value=preco_emb_val if preco_emb_val is not None else "")
        c_preco.alignment = Alignment(horizontal="right", vertical="center")
        c_preco.number_format = "R$ #,##0.00"

        # Aplica bordas e cores
        for col_i in range(1, 10):
            cell = ws.cell(row=current_row, column=col_i)
            cell.border = BORDER_THIN
            if col_i in (5, 6, 7, 8, 9):
                # Destaca suavemente colunas editáveis
                cell.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
            else:
                cell.fill = fill

        ws.row_dimensions[current_row].height = 22
        current_row += 1

    # Salva o arquivo em memória
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    b64_content = base64.b64encode(buffer.read()).decode("utf-8")

    import re
    nome_arquivo = f"cotacao_rodada_{id_rodada}"
    if fornecedor_nome:
        safe_forn = re.sub(r'[^\w\s-]', '', fornecedor_nome).strip().replace(' ', '_').lower()
        nome_arquivo += f"_{safe_forn}"
    nome_arquivo += ".xlsx"

    return {
        "sucesso": True,
        "nome_arquivo": nome_arquivo,
        "conteudo_base64": b64_content,
        "total_itens": len(necessidades),
    }


def process_quote_excel(
    conn: sqlite3.Connection,
    id_rodada: int,
    id_fornecedor: int,
    conteudo_base64: str,
) -> Dict[str, Any]:
    """
    Processa um arquivo Excel (.xlsx) enviado via base64, lê as colunas de cotação,
    auto-cadastra novos produtos caso não existam e insere/atualiza os registros em cotacoes.
    """
    raw_data = base64.b64decode(conteudo_base64)
    buffer = io.BytesIO(raw_data)
    wb = openpyxl.load_workbook(buffer, data_only=True)
    ws = wb.active

    cursor = conn.cursor()

    # Valida fornecedor e rodada
    cursor.execute("SELECT id, nome FROM fornecedores WHERE id = ?", (id_fornecedor,))
    forn_row = cursor.fetchone()
    if not forn_row:
        raise ValueError(f"Fornecedor #{id_fornecedor} não encontrado.")

    cursor.execute("SELECT id, descricao FROM rodadas WHERE id = ?", (id_rodada,))
    if not cursor.fetchone():
        raise ValueError(f"Rodada #{id_rodada} não encontrada.")

    # Mapa de produtos existentes para busca rápida
    cursor.execute("SELECT id, nome FROM produtos")
    produtos_por_id = {r["id"]: r["nome"] for r in cursor.fetchall()}
    produtos_por_nome = {r["nome"].strip().lower(): r["id"] for r in cursor.fetchall()}

    importados = 0
    ignorados = 0
    produtos_criados = 0
    erros: List[str] = []

    # Procura a linha de cabeçalho
    header_row_idx = 4
    for r in range(1, 10):
        val_a = ws.cell(row=r, column=1).value
        val_b = ws.cell(row=r, column=2).value
        if val_a and ("id" in str(val_a).lower() or "produto" in str(val_b).lower()):
            header_row_idx = r
            break

    # Itera pelas linhas de dados
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        id_prod_val = ws.cell(row=row_idx, column=1).value
        nome_prod_val = ws.cell(row=row_idx, column=2).value

        # Pode ser modelo novo (9 colunas: Marca na col 5, Embalagem na 6, Qtd na 7, Unidade na 8, Preço na 9)
        # ou modelo anterior (8 colunas)
        if ws.max_column >= 9:
            marca_val = ws.cell(row=row_idx, column=5).value
            embalagem_val = ws.cell(row=row_idx, column=6).value
            qtd_emb_val = ws.cell(row=row_idx, column=7).value
            unidade_val = ws.cell(row=row_idx, column=8).value
            preco_emb_val = ws.cell(row=row_idx, column=9).value
        else:
            marca_val = None
            embalagem_val = ws.cell(row=row_idx, column=6).value
            qtd_emb_val = ws.cell(row=row_idx, column=7).value
            unidade_val = "UN"
            preco_emb_val = ws.cell(row=row_idx, column=8).value

        if not id_prod_val and not nome_prod_val:
            continue

        # Valida se há preço preenchido antes de qualquer ação
        if preco_emb_val is None or str(preco_emb_val).strip() == "":
            ignorados += 1
            continue

        try:
            preco_emb = float(preco_emb_val)
        except (ValueError, TypeError):
            erros.append(f"Linha {row_idx}: Preço inválido '{preco_emb_val}'.")
            ignorados += 1
            continue

        if preco_emb <= 0:
            ignorados += 1
            continue

        # Identifica produto existente por ID ou nome
        id_produto = None
        if id_prod_val:
            try:
                id_tentativa = int(id_prod_val)
                if id_tentativa in produtos_por_id:
                    id_produto = id_tentativa
            except (ValueError, TypeError):
                pass

        if not id_produto and nome_prod_val:
            nome_norm = str(nome_prod_val).strip().lower()
            id_produto = produtos_por_nome.get(nome_norm)

        # Auto-cadastro de produto caso não exista e haja preço cotado
        if not id_produto and nome_prod_val and str(nome_prod_val).strip():
            nome_novo = str(nome_prod_val).strip()
            cursor.execute(
                "INSERT INTO produtos (nome, categoria, ativo) VALUES (?, NULL, 1)",
                (nome_novo,),
            )
            id_produto = cursor.lastrowid
            produtos_por_id[id_produto] = nome_novo
            produtos_por_nome[nome_novo.lower()] = id_produto
            produtos_criados += 1

        if not id_produto:
            erros.append(f"Linha {row_idx}: Nome do produto não informado.")
            ignorados += 1
            continue

        # Garante que o produto conste nas necessidades da rodada
        cursor.execute(
            """
            INSERT INTO necessidades (id_rodada, id_produto)
            VALUES (?, ?)
            ON CONFLICT(id_rodada, id_produto) DO NOTHING
            """,
            (id_rodada, id_produto),
        )

        # Valida quantidade por embalagem
        try:
            qtd_emb = float(qtd_emb_val) if qtd_emb_val else 1.0
            if qtd_emb <= 0:
                qtd_emb = 1.0
        except (ValueError, TypeError):
            qtd_emb = 1.0

        marca_desc = str(marca_val).strip() if marca_val and str(marca_val).strip() else None
        embalagem_desc = str(embalagem_val).strip() if embalagem_val else "Unidade"
        unidade_desc = str(unidade_val).strip().upper() if unidade_val else "UN"

        # Insere ou atualiza na tabela de cotações
        cursor.execute(
            """
            INSERT INTO cotacoes (
                id_rodada, id_fornecedor, id_produto,
                marca, embalagem, qtd_por_embalagem, unidade, preco_embalagem
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(id_rodada, id_fornecedor, id_produto) DO UPDATE SET
                marca = excluded.marca,
                embalagem = excluded.embalagem,
                qtd_por_embalagem = excluded.qtd_por_embalagem,
                unidade = excluded.unidade,
                preco_embalagem = excluded.preco_embalagem
            """,
            (id_rodada, id_fornecedor, id_produto, marca_desc, embalagem_desc, qtd_emb, unidade_desc, preco_emb),
        )
        importados += 1

    conn.commit()

    return {
        "sucesso": True,
        "importados": importados,
        "produtos_criados": produtos_criados,
        "ignorados": ignorados,
        "erros": erros,
        "fornecedor_nome": forn_row["nome"],
    }


def export_products_excel_db(conn: sqlite3.Connection) -> Dict[str, Any]:
    """Exporta o catálogo mestre completo de produtos para Excel."""
    cursor = conn.cursor()
    cursor.execute("SELECT id, nome, categoria FROM produtos ORDER BY categoria ASC, nome ASC")
    produtos = [dict(r) for r in cursor.fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.views.sheetView[0].showGridLines = True

    headers = [("ID", 8), ("Nome do Produto", 40), ("Categoria", 25)]
    ws.row_dimensions[1].height = 24

    for col_i, (h_text, width) in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_i, value=h_text)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(col_i)].width = width

    for row_i, p in enumerate(produtos, 2):
        ws.cell(row=row_i, column=1, value=p["id"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_i, column=2, value=p["nome"]).alignment = Alignment(horizontal="left")
        ws.cell(row=row_i, column=3, value=p["categoria"] or "-").alignment = Alignment(horizontal="center")
        for c in range(1, 4):
            ws.cell(row=row_i, column=c).border = BORDER_THIN

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    b64 = base64.b64encode(buffer.read()).decode("utf-8")

    return {"sucesso": True, "nome_arquivo": "catalogo_produtos.xlsx", "conteudo_base64": b64, "total": len(produtos)}


def import_products_excel_db(conn: sqlite3.Connection, conteudo_base64: str) -> Dict[str, Any]:
    """Importa produtos em lote a partir de uma planilha Excel."""
    raw = base64.b64decode(conteudo_base64)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active

    cursor = conn.cursor()
    importados = 0
    ignorados = 0
    erros: List[str] = []

    for row_idx in range(2, ws.max_row + 1):
        nome_val = ws.cell(row=row_idx, column=2).value
        # Caso a planilha não tenha coluna ID e comece direto no nome
        if not nome_val:
            nome_val = ws.cell(row=row_idx, column=1).value
            cat_val = ws.cell(row=row_idx, column=2).value
        else:
            cat_val = ws.cell(row=row_idx, column=3).value

        if not nome_val or str(nome_val).strip() == "":
            continue

        nome = str(nome_val).strip()
        cat = str(cat_val).strip() if cat_val and str(cat_val).strip() != "-" else None

        try:
            cursor.execute(
                """
                INSERT INTO produtos (nome, categoria, ativo)
                VALUES (?, ?, 1)
                ON CONFLICT(nome) DO UPDATE SET
                    categoria = excluded.categoria,
                    ativo = 1
                """,
                (nome, cat),
            )
            importados += 1
        except Exception as e:
            erros.append(f"Linha {row_idx}: {e}")
            ignorados += 1

    conn.commit()
    return {"sucesso": True, "importados": importados, "ignorados": ignorados, "erros": erros}


def export_suppliers_excel_db(conn: sqlite3.Connection) -> Dict[str, Any]:
    """Exporta o catálogo completo de fornecedores para uma planilha Excel estilizada."""
    cursor = conn.cursor()
    cursor.execute("SELECT id, nome, contato, telefone, email, pedido_minimo FROM fornecedores ORDER BY nome ASC")
    fornecedores = [dict(r) for r in cursor.fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fornecedores"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        ("ID", 8),
        ("Fornecedor / Razão Social", 35),
        ("Contato", 22),
        ("Telefone / WhatsApp", 20),
        ("E-mail", 30),
        ("Pedido Mínimo (R$)", 20),
    ]
    ws.row_dimensions[1].height = 24

    for col_i, (h_text, width) in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_i, value=h_text)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_THIN
        ws.column_dimensions[get_column_letter(col_i)].width = width

    for row_i, f in enumerate(fornecedores, 2):
        ws.cell(row=row_i, column=1, value=f["id"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_i, column=2, value=f["nome"]).alignment = Alignment(horizontal="left")
        ws.cell(row=row_i, column=3, value=f["contato"] or "-").alignment = Alignment(horizontal="left")
        ws.cell(row=row_i, column=4, value=f["telefone"] or "-").alignment = Alignment(horizontal="center")
        ws.cell(row=row_i, column=5, value=f["email"] or "-").alignment = Alignment(horizontal="left")
        c_min = ws.cell(row=row_i, column=6, value=f["pedido_minimo"])
        c_min.number_format = 'R$ #,##0.00'
        c_min.alignment = Alignment(horizontal="right")
        for c in range(1, 7):
            ws.cell(row=row_i, column=c).border = BORDER_THIN

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    b64 = base64.b64encode(buffer.read()).decode("utf-8")

    return {
        "sucesso": True,
        "nome_arquivo": "catalogo_fornecedores.xlsx",
        "conteudo_base64": b64,
        "total": len(fornecedores),
    }


def import_suppliers_excel_db(conn: sqlite3.Connection, conteudo_base64: str) -> Dict[str, Any]:
    """Importa fornecedores em lote a partir de uma planilha Excel."""
    raw = base64.b64decode(conteudo_base64)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active

    cursor = conn.cursor()
    importados = 0
    ignorados = 0
    erros: List[str] = []

    for row_idx in range(2, ws.max_row + 1):
        nome_val = ws.cell(row=row_idx, column=2).value
        if not nome_val:
            nome_val = ws.cell(row=row_idx, column=1).value
            contato_val = ws.cell(row=row_idx, column=2).value
            tel_val = ws.cell(row=row_idx, column=3).value
            email_val = ws.cell(row=row_idx, column=4).value
            min_val = ws.cell(row=row_idx, column=5).value
        else:
            contato_val = ws.cell(row=row_idx, column=3).value
            tel_val = ws.cell(row=row_idx, column=4).value
            email_val = ws.cell(row=row_idx, column=5).value
            min_val = ws.cell(row=row_idx, column=6).value

        if not nome_val or str(nome_val).strip() == "":
            continue

        nome = str(nome_val).strip()
        contato = str(contato_val).strip() if contato_val and str(contato_val).strip() != "-" else None
        tel = str(tel_val).strip() if tel_val and str(tel_val).strip() != "-" else None
        email = str(email_val).strip() if email_val and str(email_val).strip() != "-" else None

        try:
            ped_min = float(min_val) if min_val else 0.0
        except (ValueError, TypeError):
            ped_min = 0.0

        try:
            cursor.execute(
                """
                INSERT INTO fornecedores (nome, contato, telefone, email, pedido_minimo, ativo)
                VALUES (?, ?, ?, ?, ?, 1)
                ON CONFLICT(nome) DO UPDATE SET
                    contato = excluded.contato,
                    telefone = excluded.telefone,
                    email = excluded.email,
                    pedido_minimo = excluded.pedido_minimo,
                    ativo = 1
                """,
                (nome, contato, tel, email, ped_min),
            )
            importados += 1
        except Exception as e:
            erros.append(f"Linha {row_idx}: {e}")
            ignorados += 1

    conn.commit()
    return {"sucesso": True, "importados": importados, "ignorados": ignorados, "erros": erros}
