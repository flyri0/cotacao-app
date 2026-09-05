import base64
import io
import sqlite3
import unittest

import openpyxl

from backend.api import Api
from backend.db import create_schema, seed_data
from backend.services.excel_service import (
    generate_quote_template_excel,
    import_products_excel_db,
    import_suppliers_excel_db,
    process_quote_excel,
)


class TestExcelService(unittest.TestCase):
    def setUp(self) -> None:
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys = ON;")
        create_schema(self.conn)
        seed_data(self.conn)

        self.api = Api(db_path=":memory:")
        self.api._get_connection = lambda: self.conn

    def tearDown(self) -> None:
        self.conn.close()

    def test_gerar_e_processar_planilha_cotacao(self) -> None:
        # Gera modelo de cotação para rodada 4 (aberta)
        res = self.api.export_quote_spreadsheet(id_rodada=4, id_fornecedor=1)
        self.assertTrue(res["sucesso"])
        self.assertIn("conteudo_base64", res)
        self.assertGreater(res["total_itens"], 0)

        # Lê o Excel gerado e preenche preços para simular fornecedor
        raw = base64.b64decode(res["conteudo_base64"])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb.active

        # Preenche preços na linha 5 e 6
        ws.cell(row=5, column=5, value="Marca Alpha")  # Marca
        ws.cell(row=5, column=6, value="Caixa c/ 10")  # Embalagem
        ws.cell(row=5, column=7, value=10)  # Qtd
        ws.cell(row=5, column=8, value="CX")  # Unidade
        ws.cell(row=5, column=9, value=50.0)  # Preço da embalagem (R$ 5.00 unitário)

        ws.cell(row=6, column=5, value="Marca Beta")
        ws.cell(row=6, column=6, value="Fardo c/ 20")
        ws.cell(row=6, column=7, value=20)
        ws.cell(row=6, column=8, value="PCT")
        ws.cell(row=6, column=9, value=160.0)

        # Adiciona uma nova linha com produto inédito na planilha para testar auto-cadastro
        ws.cell(row=7, column=1, value="")
        ws.cell(row=7, column=2, value="Produto Importado Automaticamente")
        ws.cell(row=7, column=5, value="Marca Gamma")
        ws.cell(row=7, column=6, value="Galão 5L")
        ws.cell(row=7, column=7, value=5)
        ws.cell(row=7, column=8, value="L")
        ws.cell(row=7, column=9, value=45.0)

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        filled_b64 = base64.b64encode(out_buf.read()).decode("utf-8")

        # Importa de volta via API
        import_res = self.api.import_quote_spreadsheet(id_rodada=4, id_fornecedor=1, conteudo_base64=filled_b64)
        self.assertTrue(import_res["sucesso"])
        self.assertGreaterEqual(import_res["importados"], 3)
        self.assertEqual(import_res.get("produtos_criados"), 1)

    def test_exportar_e_importar_produtos_excel(self) -> None:
        exp = self.api.export_products_excel()
        self.assertTrue(exp["sucesso"])
        self.assertEqual(exp["total"], 50)

        # Cria uma planilha com novo produto
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Nome do Produto", "Categoria"])
        ws.append(["", "Produto Teste Excel Novo", "Testes"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        novo_b64 = base64.b64encode(buf.read()).decode("utf-8")

        imp = self.api.import_products_excel(novo_b64)
        self.assertTrue(imp["sucesso"])
        self.assertEqual(imp["importados"], 1)

        produtos = self.api.list_products()
        self.assertEqual(len(produtos), 51)

    def test_importar_fornecedores_excel(self) -> None:
        # Cria uma planilha com novo fornecedor
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Fornecedor / Razão Social", "Contato", "Telefone", "E-mail", "Pedido Mínimo (R$)"])
        ws.append(["", "Fornecedor Excel Ltda", "Ana", "11988887777", "ana@excel.com", 600.0])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        novo_b64 = base64.b64encode(buf.read()).decode("utf-8")

        imp = self.api.import_suppliers_excel(novo_b64)
        self.assertTrue(imp["sucesso"])
        self.assertEqual(imp["importados"], 1)

        fornecedores = self.api.list_suppliers()
        self.assertEqual(len(fornecedores), 9)

    def test_generate_quote_template_excel_errors_and_options(self) -> None:
        """Testa validações na geração de template de cotação."""
        # Rodada inexistente -> ValueError
        with self.assertRaises(ValueError):
            generate_quote_template_excel(self.conn, id_rodada=99999)

        # Geração sem id_fornecedor
        res_sem_forn = generate_quote_template_excel(self.conn, id_rodada=1, id_fornecedor=None)
        self.assertTrue(res_sem_forn["sucesso"])
        self.assertTrue(res_sem_forn["nome_arquivo"].endswith(".xlsx"))
        self.assertIn("Janeiro_2026", res_sem_forn["nome_arquivo"])

    def test_process_quote_excel_validations_and_legacy_columns(self) -> None:
        """Testa importação de cotações com planilha legado de 8 colunas e validações."""
        # Planilha modelo legado (8 colunas: sem coluna Marca)
        wb = openpyxl.Workbook()
        ws = wb.active
        # Cabeçalho na linha 4
        ws.cell(row=4, column=1, value="ID")
        ws.cell(row=4, column=2, value="Produto")
        ws.cell(row=4, column=3, value="Categoria")
        ws.cell(row=4, column=4, value="Ref")
        ws.cell(row=4, column=5, value="Info")
        ws.cell(row=4, column=6, value="Embalagem")
        ws.cell(row=4, column=7, value="Qtd")
        ws.cell(row=4, column=8, value="Preço Embalagem")

        # Linha 5: Preço válido com id não numérico e qtd negativa (testa fallback para 1.0)
        ws.cell(row=5, column=1, value="abc_invalido")
        ws.cell(row=5, column=2, value="Detergente Líquido Neutro 500ml")
        ws.cell(row=5, column=6, value="Caixa c/ 10")
        ws.cell(row=5, column=7, value=-2)  # qtd <= 0 -> fallback para 1.0
        ws.cell(row=5, column=8, value=30.0)

        # Linha 6: Preço vazio -> ignorado
        ws.cell(row=6, column=1, value=1)
        ws.cell(row=6, column=2, value="Detergente Líquido Neutro 500ml")
        ws.cell(row=6, column=8, value="")

        # Linha 7: Preço <= 0 -> ignorado
        ws.cell(row=7, column=1, value=1)
        ws.cell(row=7, column=2, value="Detergente Líquido Neutro 500ml")
        ws.cell(row=7, column=8, value=0.0)

        # Linha 8: Preço inválido não numérico -> erro registrado
        ws.cell(row=8, column=1, value=1)
        ws.cell(row=8, column=2, value="Detergente Líquido Neutro 500ml")
        ws.cell(row=8, column=8, value="preço_invalido_texto")

        # Linha 9: Preço válido mas sem id e sem nome -> erro de produto não informado
        ws.cell(row=9, column=1, value="")
        ws.cell(row=9, column=2, value="")
        ws.cell(row=9, column=8, value=25.0)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        b64_legado = base64.b64encode(buf.read()).decode("utf-8")

        # Rodada fechada com planilha válida -> ValueError
        with self.assertRaises(ValueError):
            process_quote_excel(self.conn, id_rodada=1, id_fornecedor=1, conteudo_base64=b64_legado)

        # Rodada inexistente com planilha válida
        with self.assertRaises(ValueError):
            process_quote_excel(self.conn, id_rodada=99999, id_fornecedor=1, conteudo_base64=b64_legado)

        # Fornecedor inexistente com planilha válida
        with self.assertRaises(ValueError):
            process_quote_excel(self.conn, id_rodada=4, id_fornecedor=99999, conteudo_base64=b64_legado)

        res_legado = process_quote_excel(self.conn, id_rodada=4, id_fornecedor=1, conteudo_base64=b64_legado)
        self.assertTrue(res_legado["sucesso"])
        self.assertGreaterEqual(res_legado["importados"], 1)
        self.assertGreater(res_legado["ignorados"], 0)
        self.assertGreater(len(res_legado["erros"]), 0)

    def test_import_products_excel_edge_cases(self) -> None:
        """Testa casos de borda e duplicatas na importação de produtos."""
        from unittest.mock import MagicMock

        # 1. Planilha vazia sem linhas de dados
        wb_vazio = openpyxl.Workbook()
        buf = io.BytesIO()
        wb_vazio.save(buf)
        buf.seek(0)
        b64_vazio = base64.b64encode(buf.read()).decode("utf-8")
        res_vazio = import_products_excel_db(self.conn, b64_vazio)
        self.assertEqual(res_vazio["importados"], 0)

        # 2. Planilha com produto com ID, sem ID (linhas 439-440), linha vazia (linha 445)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Nome do Produto", "Categoria"])
        ws.append(["", "Detergente Líquido Neutro 500ml", "Limpeza"])
        ws.append(["Produto Teste Fallback Col1", "", "Geral"])  # Coluna 2 vazia -> pega col 1
        ws.append(["", "   ", "Sem Nome"])  # Linha com nome vazio (linha 445)

        buf2 = io.BytesIO()
        wb.save(buf2)
        buf2.seek(0)
        b64_prod = base64.b64encode(buf2.read()).decode("utf-8")

        res_prod = import_products_excel_db(self.conn, b64_prod)
        self.assertTrue(res_prod["sucesso"])
        self.assertEqual(res_prod["importados"], 2)

        # 3. Exceção simulada no cursor.execute (linhas 462-464)
        mock_conn = MagicMock()
        mock_conn.cursor.return_value.execute.side_effect = sqlite3.DatabaseError("Erro simulado")
        res_err = import_products_excel_db(mock_conn, b64_prod)
        self.assertEqual(res_err["importados"], 0)
        self.assertGreater(res_err["ignorados"], 0)
        self.assertGreater(len(res_err["erros"]), 0)

    def test_import_suppliers_excel_edge_cases(self) -> None:
        """Testa casos de borda, duplicatas e parsing de pedido mínimo em fornecedores."""
        from unittest.mock import MagicMock

        # 1. Planilha vazia
        wb_vazio = openpyxl.Workbook()
        buf = io.BytesIO()
        wb_vazio.save(buf)
        buf.seek(0)
        b64_vazio = base64.b64encode(buf.read()).decode("utf-8")
        res_vazio = import_suppliers_excel_db(self.conn, b64_vazio)
        self.assertEqual(res_vazio["importados"], 0)

        # 2. Fornecedor na coluna 1 (sem coluna ID, linhas 538-542), linha sem nome e formatos
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Fornecedor", "Contato", "Telefone", "E-mail", "Pedido Mínimo"])
        ws.append(["", "Distribuidora São Paulo", "João", "", "", "500"])
        ws.append(["Fornecedor Fallback Col1", "", "", "", "", ""])  # Coluna 2 vazia -> pega col 1
        ws.append(["", "   ", "", "", "", ""])  # Nome vazio (linha 550)
        ws.append(["", "Novo Fornecedor Pedido BRL", "Maria", "", "", "R$ 1.500,50"])  # BRL float
        ws.append(["", "Novo Fornecedor Pedido Invalido", "José", "", "", "invalido_texto"])  # Inv -> 0.0

        buf2 = io.BytesIO()
        wb.save(buf2)
        buf2.seek(0)
        b64_forn = base64.b64encode(buf2.read()).decode("utf-8")

        res_forn = import_suppliers_excel_db(self.conn, b64_forn)
        self.assertTrue(res_forn["sucesso"])
        self.assertEqual(res_forn["importados"], 4)

        # 3. Exceção simulada no cursor.execute (linhas 577-579)
        mock_conn = MagicMock()
        mock_conn.cursor.return_value.execute.side_effect = sqlite3.DatabaseError("Erro fornecedor")
        res_err = import_suppliers_excel_db(mock_conn, b64_forn)
        self.assertEqual(res_err["importados"], 0)
        self.assertGreater(res_err["ignorados"], 0)
        self.assertGreater(len(res_err["erros"]), 0)
